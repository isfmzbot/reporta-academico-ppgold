import streamlit as st
import xml.etree.ElementTree as ET
import pandas as pd
import io
import datetime
import unicodedata
import re
import PyPDF2
import docx
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==========================================
# 1. FUNÇÕES DE LIMPEZA E EXTRAÇÃO (LATTES)
# ==========================================
def normalizar_nome(nome):
    if not isinstance(nome, str) or pd.isna(nome):
        return ""
    nome = unicodedata.normalize('NFKD', nome).encode('ASCII', 'ignore').decode('utf-8')
    return nome.upper().strip()

def extrair_dados_lattes(arquivo_xml):
    tree = ET.parse(arquivo_xml)
    root = tree.getroot()
    dados_gerais = root.find('DADOS-GERAIS')
    nome_pesquisador = dados_gerais.get('NOME-COMPLETO') if dados_gerais is not None else "Não Identificado"
    
    artigos_lista = []
    producao = root.find('PRODUCAO-BIBLIOGRAFICA')
    if producao is not None:
        artigos = producao.find('ARTIGOS-PUBLICADOS')
        if artigos is not None:
            for artigo in artigos.findall('ARTIGO-PUBLICADO'):
                dados_basicos = artigo.find('DADOS-BASICOS-DO-ARTIGO')
                detalhamento = artigo.find('DETALHAMENTO-DO-ARTIGO')
                
                if dados_basicos is not None:
                    pais = dados_basicos.get('PAIS-DE-PUBLICACAO', 'Não Informado')
                    if pais == 'Não Informado' or pais.strip() == '': 
                        abrangencia = 'Não Informado'
                    elif pais.upper() == 'BRASIL': 
                        abrangencia = 'Nacional'
                    else: 
                        abrangencia = 'Internacional'

                    issn = "Ausente"
                    revista = "Ausente"
                    if detalhamento is not None:
                        issn_cru = detalhamento.get('ISSN', 'Ausente')
                        revista = detalhamento.get('TITULO-DO-PERIODICO-OU-REVISTA', 'Ausente')
                        
                        if issn_cru != 'Ausente' and len(issn_cru) == 8: 
                            issn = f"{issn_cru[:4]}-{issn_cru[4:]}"
                        else: 
                            issn = issn_cru

                    artigos_lista.append({
                        'Professor': nome_pesquisador,
                        'Titulo': dados_basicos.get('TITULO-DO-ARTIGO'),
                        'Ano': dados_basicos.get('ANO-DO-ARTIGO'),
                        'Revista': revista,
                        'Revista_Norm': normalizar_nome(revista),
                        'ISSN': issn,
                        'País': pais,
                        'Abrangência': abrangencia
                    })
    return artigos_lista


# ==========================================
# 2. MOTOR DE EXTRAÇÃO DE ARQUIVOS PTT (PDF/WORD)
# ==========================================
def limpar_cabecalhos_texto(texto):
    padroes_remover = [
        r'Template:\s*Artigo tecnológico.*?\n',
        r'ESCOLA DE\s*ADMINISTRA[CÇ][ÃA]O\s*PPGOLD.*?\n',
        r'Universidade Federal do Paraná.*?\n',
        r'Setor de Ciências Sociais Aplicadas.*?\n',
        r'Programa de Pós-Graduação em.*?\n',
        r'Gestão de Organizações, Liderança e Decisão.*?\n',
        r'Mestrado Profissional.*?\n',
        r'UFPR\s*\n'
    ]
    texto_limpo = texto
    for p in padroes_remover:
        texto_limpo = re.sub(p, '', texto_limpo, flags=re.IGNORECASE)
    return texto_limpo

def extrair_campo_entre(texto, inicio_regex, fim_regex):
    padrao = rf'{inicio_regex}\s*(.*?)\s*(?={fim_regex})'
    match = re.search(padrao, texto, re.IGNORECASE | re.DOTALL)
    if match:
        valor = re.sub(r'\s+', ' ', match.group(1)).strip()
        return valor if valor else "Não Preenchido"
    return "Não Encontrado"

def extrair_alternativa_criterio(texto, ancora_inicio, ancora_fim):
    padrao_bloco = rf'{ancora_inicio}(.*?)(?={ancora_fim})'
    match_bloco = re.search(padrao_bloco, texto, re.IGNORECASE | re.DOTALL)
    if not match_bloco:
        return "Bloco não localizado"
    
    bloco = match_bloco.group(1)
    padrao_marcacao = r'(?:\(\s*[xX☑■✓✔]\s*\)|\[\s*[xX☑■✓✔]\s*\]|[☑■✓✔]|(?:^|\n)\s*[xX]\s*[-.)]?)\s*(Ausente|M[ée]dia|M[ée]dio|Baixa|Baixo|Alta|Alto)'
    match_op = re.search(padrao_marcacao, bloco, re.IGNORECASE)
    if match_op:
        return match_op.group(1).capitalize()
    return "Não Assinalado"

def extrair_tipo_trabalho(texto):
    padrao_bloco = r'Tipo de trabalho desenvolvido:(.*?)(?=CRITÉRIOS DE AVALIAÇÃO)'
    match_bloco = re.search(padrao_bloco, texto, re.IGNORECASE | re.DOTALL)
    bloco = match_bloco.group(1) if match_bloco else texto
    
    opcoes_conhecidas = [
        "Base de Dados Técnico-científica",
        "Curso para formação profissional",
        "Empresa ou organização social inovadora",
        "Material didático",
        "Norma ou marco regulatório",
        "Registro de Patente",
        "Processo/tecnologia e/ou produtos não patenteáveis",
        "Produto bibliográfico na forma de artigo tecnológico",
        "Produtos e processos em sigilo",
        "Relatório técnico conclusivo (consultorias e afins)",
        "Software ou aplicativo",
        "Tecnologia social",
        "Outro"
    ]
    
    for op in opcoes_conhecidas:
        op_escapada = re.escape(op)
        padrao = rf'(?:\(\s*[xX☑■✓✔]\s*\)|\[\s*[xX☑■✓✔]\s*\]|[☑■✓✔]|\b[xX]\b)\s*\|?\s*{op_escapada}'
        if re.search(padrao, bloco, re.IGNORECASE):
            if op == "Outro":
                m_outro = re.search(r'Outro:\s*(?:especificar claramente)?\s*(.*)', bloco, re.IGNORECASE)
                detalhe = m_outro.group(1).strip() if m_outro and m_outro.group(1).strip() else "Outro"
                return f"Outro ({detalhe})"
            return op
            
    return "Não Identificado"

def ler_texto_docx_completo(arquivo):
    doc = docx.Document(arquivo)
    linhas = []
    for elemento in doc.element.body:
        if elemento.tag.endswith('p'):
            textos = [t.text for t in elemento.iter() if t.tag.endswith('t') and t.text]
            linhas.append("".join(textos))
        elif elemento.tag.endswith('tbl'):
            for row in elemento.iter():
                if row.tag.endswith('tr'):
                    celulas = []
                    for cell in row:
                        if cell.tag.endswith('tc'):
                            t_cel = [t.text for t in cell.iter() if t.tag.endswith('t') and t.text]
                            celulas.append("".join(t_cel).strip())
                    linhas.append(" | ".join(celulas))
    return "\n".join(linhas)

def extrair_dados_arquivo_ptt(arquivo):
    try:
        extensao = arquivo.name.split('.')[-1].lower()
        texto_bruto = ""
        
        if extensao == 'pdf':
            leitor = PyPDF2.PdfReader(arquivo)
            for pagina in leitor.pages:
                texto_bruto += (pagina.extract_text() or "") + "\n"
        elif extensao in ['docx', 'doc']:
            texto_bruto = ler_texto_docx_completo(arquivo)
        else:
            return {"Título do trabalho": "Formato não suportado", "Arquivo Origem": arquivo.name}

        texto = limpar_cabecalhos_texto(texto_bruto)
        
        titulo = extrair_campo_entre(texto, r'T[íi]tulo do\s*trabalho:', r'Autor do trabalho:')
        autor = extrair_campo_entre(texto, r'Autor do trabalho:', r'Orientador respons[áa]vel:')
        orientador = extrair_campo_entre(texto, r'Orientador respons[áa]vel:', r'Per[íi]odo de desenvolvimento')
        periodo = extrair_campo_entre(texto, r'Per[íi]odo de desenvolvimento do trabalho:\s*(?:\(m[êe]s in[íi]cio;\s*m[êe]s de fim/ano\):)?', r'Organiza[çc][ãa]o em que o trabalho foi aplicado:')
        organizacao = extrair_campo_entre(texto, r'Organiza[çc][ãa]o em que o trabalho foi aplicado:', r'Resumo do trabalho')
        resumo = extrair_campo_entre(texto, r'Resumo do trabalho em at[ée] \d+ palavras:', r'Tipo de trabalho desenvolvido:')
        
        tipo_trabalho = extrair_tipo_trabalho(texto)
        
        imp_realizado = extrair_alternativa_criterio(texto, r'Impacto realizado', r'Impacto potencial')
        imp_potencial = extrair_alternativa_criterio(texto, r'Impacto potencial', r'Justificativa|2\.\s*APLICABILIDADE|APLICABILIDADE')
        aplic_realizada = extrair_alternativa_criterio(texto, r'Aplicabilidade realizada', r'Aplicabilidade potencial')
        aplic_potencial = extrair_alternativa_criterio(texto, r'Aplicabilidade potencial', r'Replicabilidade')
        replicabilidade = extrair_alternativa_criterio(texto, r'Replicabilidade', r'Justificativa|3\.\s*INOVA[ÇC][ÃA]O|INOVA[ÇC][ÃA]O')
        inovacao = extrair_alternativa_criterio(texto, r'(?:3\.\s*)?INOVA[ÇC][ÃA]O', r'Justificativa|4\.\s*COMPLEXIDADE|COMPLEXIDADE')
        complexidade = extrair_alternativa_criterio(texto, r'(?:4\.\s*)?COMPLEXIDADE', r'Justificativa|5\.\s*RELEV[ÂA]NCIA|RELEV[ÂA]NCIA')
        
        return {
            "Título do trabalho": titulo,
            "Autor do trabalho": autor,
            "Orientador responsável": orientador,
            "Período de desenvolvimento": periodo,
            "Organização em que o trabalho foi aplicado": organizacao,
            "Resumo do trabalho (até 25 palavras)": resumo,
            "Tipo de trabalho desenvolvido": tipo_trabalho,
            "Impacto realizado": imp_realizado,
            "Impacto potencial": imp_potencial,
            "Aplicabilidade realizada": aplic_realizada,
            "Aplicabilidade potencial": aplic_potencial,
            "Replicabilidade": replicabilidade,
            "Inovação": inovacao,
            "Complexidade": complexidade,
            "Arquivo Origem": arquivo.name
        }
    except Exception as e:
        return {
            "Título do trabalho": f"Erro na leitura: {e}",
            "Autor do trabalho": "Erro",
            "Orientador responsável": "",
            "Período de desenvolvimento": "",
            "Organização em que o trabalho foi aplicado": "",
            "Resumo do trabalho (até 25 palavras)": "",
            "Tipo de trabalho desenvolvido": "",
            "Impacto realizado": "",
            "Impacto potencial": "",
            "Aplicabilidade realizada": "",
            "Aplicabilidade potencial": "",
            "Replicabilidade": "",
            "Inovação": "",
            "Complexidade": "",
            "Arquivo Origem": arquivo.name
        }


# ==========================================
# 3. MOTOR DE CÁLCULO - REGRAS ÁREA 27
# ==========================================
def calcular_nota_area27(row, dict_bases, bases_carregadas):
    issn = row['ISSN']
    revista_norm = row['Revista_Norm']
    abrangencia_original = row['Abrangência']
    
    abdc = dict_bases['abdc'].get(issn, "Ausente") if "ABDC" in bases_carregadas else ""
    abs_nota = dict_bases['abs'].get(issn, "Ausente") if "ABS" in bases_carregadas else ""
    jcr = dict_bases['jcr'].get(issn, "Ausente") if "JCR" in bases_carregadas else ""
    sjr = dict_bases['sjr'].get(issn, "Ausente") if "SJR" in bases_carregadas else ""
    
    if "SPELL" in bases_carregadas:
        spell_info = dict_bases['spell'].get(revista_norm, {})
        spell_rank = spell_info.get('rank', "Ausente")
    else:
        spell_rank = ""
        
    is_scielo = False
    scielo_status = ""
    if "SciELO" in bases_carregadas:
        is_scielo = revista_norm in dict_bases['scielo']
        scielo_status = "Presente" if is_scielo else "Ausente"

    abrangencia_final = abrangencia_original
    if abrangencia_original == 'Não Informado':
        if ("SciELO" in bases_carregadas and is_scielo) or ("SPELL" in bases_carregadas and spell_rank not in ["Ausente", ""]):
            abrangencia_final = 'Nacional (Deduzido)'
        elif jcr not in ["Ausente", ""] or sjr not in ["Ausente", ""] or abdc not in ["Ausente", ""] or abs_nota not in ["Ausente", ""]:
            abrangencia_final = 'Internacional (Deduzido)'

    nota = "Não Classificado"
    if abdc in ['A*', 'A'] or str(abs_nota) in ['2', '3', '4', '4*'] or jcr == 'Q1' or sjr == 'Q1':
        nota = 'MB'
    elif abdc == 'B' or str(abs_nota) == '1' or jcr == 'Q2' or sjr == 'Q2' or (spell_rank == 'Top 10% (D1)' and is_scielo):
        nota = 'B'
    elif abdc == 'C' or jcr == 'Q3' or sjr == 'Q3' or spell_rank in ['D2 (11-20%)', 'D3 (21-30%)', 'D4 (31-40%)']:
        nota = 'R'
    elif jcr == 'Q4' or sjr == 'Q4' or spell_rank in ['D5 (41-50%)', 'D6 (51-60%)', 'D7 (61-70%)']:
        nota = 'F'

    if is_scielo:
        if nota == 'F': 
            nota = 'R'
            scielo_status = "Presente (+1 nível: F ➔ R)"
        elif nota == 'R': 
            nota = 'B'
            scielo_status = "Presente (+1 nível: R ➔ B)"
        elif nota == 'B' and spell_rank == 'Top 10% (D1)':
            scielo_status = "Presente (Obrigatório p/ B no SPELL)"
        elif nota in ['MB', 'B']:
            scielo_status = "Presente (Sem ajuste: Limite B atingido)"

    pontos = {'MB': 8, 'B': 4, 'R': 2, 'F': 1}.get(nota, 0)
    return pd.Series([abrangencia_final, abdc, abs_nota, jcr, sjr, scielo_status, spell_rank, nota, pontos])


# ==========================================
# 4. FORMATAÇÃO AVANÇADA DO EXCEL (CAPES E PTTS)
# ==========================================
def gerar_excel_capes_estilizado(df_artigos, ano_selecionado):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        nome_aba_artigos = f'Artigos {ano_selecionado}'
        df_artigos.to_excel(writer, index=False, sheet_name=nome_aba_artigos)
        
        workbook = writer.book
        ws = writer.sheets[nome_aba_artigos]
        
        ws.views.sheetView[0].showGridLines = False
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        
        cor_cabecalho = PatternFill(start_color="2A8B79", end_color="2A8B79", fill_type="solid")
        fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        cor_zebra_par = PatternFill(start_color="F2F9F7", end_color="F2F9F7", fill_type="solid")
        cor_zebra_impar = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        borda_suave = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='D3D3D3')
        )
        
        cores_conceito = {
            'MB': (PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="155724")),
            'B':  (PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="0C5460")),
            'R':  (PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="856404")),
            'F':  (PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="721C24")),
            'Não Classificado': (PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"), Font(name="Calibri", size=10, italic=True, color="383D41"))
        }
        
        colunas_esquerda = ['Autores (PPGOLD)', 'Titulo', 'Revista']
        
        ws.row_dimensions[1].height = 28
        for col_idx, cell in enumerate(ws[1], 1):
            cell.fill = cor_cabecalho
            cell.font = fonte_cabecalho
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = borda_suave
            
        for row_idx in range(2, ws.max_row + 1):
            fill_padrao = cor_zebra_par if row_idx % 2 == 0 else cor_zebra_impar
            maior_linhas_texto = 1
            
            col_conceito_idx = list(df_artigos.columns).index('Conceito Área 27') + 1 if 'Conceito Área 27' in df_artigos.columns else None
            valor_conceito = ws.cell(row=row_idx, column=col_conceito_idx).value if col_conceito_idx else None
            
            for col_idx in range(1, ws.max_column + 1):
                nome_coluna = ws.cell(row=1, column=col_idx).value
                cell = ws.cell(row=row_idx, column=col_idx)
                valor_str = str(cell.value or '')
                
                cell.border = borda_suave
                cell.fill = fill_padrao
                cell.font = Font(name="Calibri", size=11, color="212529")
                
                if nome_coluna in colunas_esquerda:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    largura_base = 45 if nome_coluna == 'Titulo' else 32
                    linhas_est = (len(valor_str) // largura_base) + 1
                    if linhas_est > maior_linhas_texto:
                        maior_linhas_texto = linhas_est
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                
                if nome_coluna in ['Conceito Área 27', 'Pontos'] and valor_conceito in cores_conceito:
                    fill_dinamico, font_dinamica = cores_conceito[valor_conceito]
                    cell.fill = fill_dinamico
                    cell.font = font_dinamica
            
            ws.row_dimensions[row_idx].height = max(22, maior_linhas_texto * 16)
            
        larguras_fixas = {
            'Autores (PPGOLD)': 32,
            'Qtd. Autores': 14,
            'Titulo': 48,
            'Ano': 10,
            'Revista': 35,
            'ISSN': 14,
            'Abrangência': 22,
            'ABDC': 10,
            'ABS': 10,
            'JCR': 10,
            'SJR': 10,
            'SciELO': 28,
            'SPELL (Classificação)': 22,
            'Conceito Área 27': 18,
            'Pontos': 12
        }
        for col_idx in range(1, ws.max_column + 1):
            nome_col = ws.cell(row=1, column=col_idx).value
            letra_col = get_column_letter(col_idx)
            ws.column_dimensions[letra_col].width = larguras_fixas.get(nome_col, 18)

        # ABA 2: EXPLICATIVO DE PONTUAÇÕES E CLASSIFICAÇÕES (ÁREA 27)
        ws_regras = workbook.create_sheet(title='Critérios e Regras')
        ws_regras.views.sheetView[0].showGridLines = False
        
        ws_regras.merge_cells('B2:F2')
        cel_titulo = ws_regras['B2']
        cel_titulo.value = "GUIA DE CLASSIFICAÇÃO E PONTUAÇÃO DE ARTIGOS - ÁREA 27 (CAPES)"
        cel_titulo.fill = cor_cabecalho
        cel_titulo.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        cel_titulo.alignment = Alignment(horizontal="center", vertical="center")
        ws_regras.row_dimensions[2].height = 32
        
        ws_regras.merge_cells('B3:F3')
        cel_sub = ws_regras['B3']
        cel_sub.value = "Critérios aplicados automaticamente pelo sistema REPORTA ACADÊMICO (PPGOLD - UFPR)"
        cel_sub.fill = PatternFill(start_color="E0F2F1", end_color="E0F2F1", fill_type="solid")
        cel_sub.font = Font(name="Calibri", size=10, italic=True, color="004D40")
        cel_sub.alignment = Alignment(horizontal="center", vertical="center")
        ws_regras.row_dimensions[3].height = 20

        cabecalhos_regras = ["Conceito", "Pontuação", "Critérios Internacionais (JCR, SJR, ABDC, ABS)", "Critérios Nacionais (SPELL e SciELO)", "Regra de Bonificação SciELO"]
        ws_regras.row_dimensions[5].height = 26
        for idx, texto_cab in enumerate(cabecalhos_regras, start=2):
            c = ws_regras.cell(row=5, column=idx, value=texto_cab)
            c.fill = PatternFill(start_color="1F6B5C", end_color="1F6B5C", fill_type="solid")
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = borda_suave

        dados_regras = [
            ("MB", 8, "JCR = Q1  ou  SJR = Q1\nABDC = A* ou A\nABS (AJG) = 2, 3, 4 ou 4*", "Não aplicável diretamente (Teto nacional = B)", "Sem ajuste adicional (Conceito máximo atingido)"),
            ("B", 4, "JCR = Q2  ou  SJR = Q2\nABDC = B\nABS (AJG) = 1", "SPELL Top 10% (Decil D1) + Indexação simultânea no SciELO", "Periódicos conceito R que estejam no SciELO sobem para B (+1 nível). Limite máximo do bônus = B."),
            ("R", 2, "JCR = Q3  ou  SJR = Q3\nABDC = C", "SPELL entre 11% e 40% (Decis D2, D3 ou D4)", "Periódicos conceito F que estejam no SciELO sobem para R (+1 nível)."),
            ("F", 1, "JCR = Q4  ou  SJR = Q4", "SPELL entre 41% e 70% (Decis D5, D6 ou D7)", "Se indexado no SciELO, é promovido automaticamente para R (+1 nível)."),
            ("Não Classificado", 0, "Ausente nas bases ou fora dos critérios acima", "SPELL abaixo de 70% (D8 a D10) ou ausente nas bases", "Sem bonificação se não atingir o critério mínimo.")
        ]

        for r_offset, (conc, pts, crit_int, crit_nac, bonus) in enumerate(dados_regras, start=6):
            ws_regras.row_dimensions[r_offset].height = 48
            fill_linha = cor_zebra_par if r_offset % 2 == 0 else cor_zebra_impar
            
            valores_linha = [conc, pts, crit_int, crit_nac, bonus]
            for c_offset, val in enumerate(valores_linha, start=2):
                cel = ws_regras.cell(row=r_offset, column=c_offset, value=val)
                cel.border = borda_suave
                
                if c_offset in [2, 3] and conc in cores_conceito:
                    f_fill, f_font = cores_conceito[conc]
                    cel.fill = f_fill
                    cel.font = f_font
                    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    cel.fill = fill_linha
                    cel.font = Font(name="Calibri", size=10, color="212529")
                    cel.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws_regras.merge_cells('B12:F12')
        cel_obs_tit = ws_regras['B12']
        cel_obs_tit.value = "NOTAS TÉCNICAS DO PROCESSAMENTO AUTOMÁTICO:"
        cel_obs_tit.font = Font(name="Calibri", size=11, bold=True, color="2A8B79")
        ws_regras.row_dimensions[12].height = 22

        notas_tecnicas = [
            "1. Prevalência do Maior Estrato: Caso o periódico possua classificação em mais de uma base (ex: Q2 no SJR e A no ABDC), o sistema atribui automaticamente o maior conceito obtido (MB).",
            "2. Dedução Inteligente de Abrangência: Quando o currículo Lattes não informa o país do periódico, o sistema deduz como 'Nacional' se estiver no SPELL/SciELO ou 'Internacional' se constar nas bases JCR/SJR/ABDC/ABS.",
            "3. Desduplicação de Coautorias: Artigos publicados em coautoria por mais de um docente do PPGOLD são agrupados em uma única linha para evitar dupla contagem, listando todos os professores autores na primeira coluna."
        ]
        for idx_nota, texto_nota in enumerate(notas_tecnicas, start=13):
            ws_regras.merge_cells(f'B{idx_nota}:F{idx_nota}')
            c_nota = ws_regras[f'B{idx_nota}']
            c_nota.value = texto_nota
            c_nota.font = Font(name="Calibri", size=10, color="333333")
            c_nota.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws_regras.row_dimensions[idx_nota].height = 22

        ws_regras.column_dimensions['A'].width = 4
        ws_regras.column_dimensions['B'].width = 18
        ws_regras.column_dimensions['C'].width = 14
        ws_regras.column_dimensions['D'].width = 36
        ws_regras.column_dimensions['E'].width = 36
        ws_regras.column_dimensions['F'].width = 38

    return buffer.getvalue()


def gerar_excel_ptts_estilizado(df_ptts):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        nome_aba = 'Fichas PTT Consolidadas'
        df_ptts.to_excel(writer, index=False, sheet_name=nome_aba)
        
        ws = writer.sheets[nome_aba]
        
        # 1. Ocultar linhas de grade
        ws.views.sheetView[0].showGridLines = False
        
        # 2. Congelar cabeçalho e ativar filtros
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        
        # Paleta de Estilos
        cor_cabecalho = PatternFill(start_color="2A8B79", end_color="2A8B79", fill_type="solid")
        fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        cor_zebra_par = PatternFill(start_color="F2F9F7", end_color="F2F9F7", fill_type="solid")
        cor_zebra_impar = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        borda_suave = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='D3D3D3')
        )
        
        # Cores dinâmicas para as respostas dos critérios avaliativos
        cores_criterios = {
            'Alta':    (PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="155724")),
            'Alto':    (PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="155724")),
            'Média':   (PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="856404")),
            'Médio':   (PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="856404")),
            'Media':   (PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="856404")),
            'Medio':   (PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="856404")),
            'Baixa':   (PatternFill(start_color="FFE8D6", end_color="FFE8D6", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="9C4100")),
            'Baixo':   (PatternFill(start_color="FFE8D6", end_color="FFE8D6", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="9C4100")),
            'Ausente': (PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"), Font(name="Calibri", size=11, bold=True, color="721C24")),
            'Não Assinalado': (PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"), Font(name="Calibri", size=10, italic=True, color="383D41")),
            'Bloco não localizado': (PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"), Font(name="Calibri", size=10, italic=True, color="383D41"))
        }
        
        colunas_criterios = [
            "Impacto realizado",
            "Impacto potencial",
            "Aplicabilidade realizada",
            "Aplicabilidade potencial",
            "Replicabilidade",
            "Inovação",
            "Complexidade"
        ]
        
        colunas_esquerda = [
            "Título do trabalho",
            "Autor do trabalho",
            "Orientador responsável",
            "Organização em que o trabalho foi aplicado",
            "Resumo do trabalho (até 25 palavras)",
            "Tipo de trabalho desenvolvido",
            "Arquivo Origem"
        ]
        
        larguras_fixas = {
            "Título do trabalho": 42,
            "Autor do trabalho": 28,
            "Orientador responsável": 28,
            "Período de desenvolvimento": 24,
            "Organização em que o trabalho foi aplicado": 32,
            "Resumo do trabalho (até 25 palavras)": 50,
            "Tipo de trabalho desenvolvido": 38,
            "Impacto realizado": 18,
            "Impacto potencial": 18,
            "Aplicabilidade realizada": 22,
            "Aplicabilidade potencial": 22,
            "Replicabilidade": 18,
            "Inovação": 16,
            "Complexidade": 16,
            "Arquivo Origem": 30
        }
        
        # 3. Estilizar Cabeçalho (Linha 1)
        ws.row_dimensions[1].height = 30
        for col_idx, cell in enumerate(ws[1], 1):
            cell.fill = cor_cabecalho
            cell.font = fonte_cabecalho
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = borda_suave
            
        # 4. Estilizar Linhas de Dados (Zebrado, Alinhamento, Wrap Text e Cores Dinâmicas)
        for row_idx in range(2, ws.max_row + 1):
            fill_padrao = cor_zebra_par if row_idx % 2 == 0 else cor_zebra_impar
            maior_linhas_texto = 1
            
            for col_idx in range(1, ws.max_column + 1):
                nome_coluna = ws.cell(row=1, column=col_idx).value
                cell = ws.cell(row=row_idx, column=col_idx)
                valor_str = str(cell.value or '').strip()
                
                cell.border = borda_suave
                cell.fill = fill_padrao
                cell.font = Font(name="Calibri", size=11, color="212529")
                
                largura_col = larguras_fixas.get(nome_coluna, 20)
                
                if nome_coluna in colunas_esquerda:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    linhas_est = (len(valor_str) // max(1, (largura_col - 4))) + 1
                    if linhas_est > maior_linhas_texto:
                        maior_linhas_texto = linhas_est
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                
                # Pintura dinâmica nas colunas de Critérios
                if nome_coluna in colunas_criterios and valor_str in cores_criterios:
                    fill_dinamico, font_dinamica = cores_criterios[valor_str]
                    cell.fill = fill_dinamico
                    cell.font = font_dinamica
            
            ws.row_dimensions[row_idx].height = max(24, maior_linhas_texto * 16)
            
        # 5. Ajuste Automático de Largura das Colunas
        for col_idx in range(1, ws.max_column + 1):
            nome_col = ws.cell(row=1, column=col_idx).value
            letra_col = get_column_letter(col_idx)
            ws.column_dimensions[letra_col].width = larguras_fixas.get(nome_col, 20)
            
    return buffer.getvalue()


# ==========================================
# 5. VALIDADOR RESTRITO DE PLANILHAS
# ==========================================
def validar_e_carregar_planilha(uploaded_file, nome_base):
    if uploaded_file is not None:
        extensao = uploaded_file.name.split('.')[-1].lower()
        if extensao not in ['xlsx', 'xls', 'csv']:
            st.error(f"Erro: O arquivo enviado para '{nome_base}' ({uploaded_file.name}) não é uma planilha válida! Apenas arquivos .xlsx, .xls ou .csv são permitidos.")
            return None
    return uploaded_file


# ==========================================
# 6. CONTROLE DE PÁGINAS E ESTÉTICA
# ==========================================
st.set_page_config(page_title="Relatórios CAPES - PPGOLD", layout="wide")

st.markdown("""
    <style>
        .header-box {
            background-color: #2a8b79;
            padding: 14px 20px;
            border-radius: 6px;
            text-align: center;
            margin-bottom: 25px;
        }
        .header-box h1 {
            color: white;
            margin: 0px;
            font-size: 26px;
            font-weight: 700;
        }
        .header-box p {
            color: #e0f2f1;
            margin: 3px 0;
            font-size: 15px;
            font-weight: 400;
        }
        .header-box h3 {
            color: white;
            margin: 3px 0 0 0;
            font-weight: 300;
            font-size: 14px;
        }

        .small-header {
            color: #2a8b79;
            font-size: 24px;
            font-weight: 800;
            text-align: center;
            margin-top: 2px;
            margin-bottom: 10px;
            letter-spacing: 1px;
        }

        div.stDownloadButton > button {
            background-color: #2a8b79 !important;
            color: #ffffff !important;
            font-size: 16px !important;
            font-weight: 700 !important;
            padding: 12px 32px !important;
            border-radius: 8px !important;
            border: none !important;
            box-shadow: 0 4px 6px rgba(0, 0, 0, 0.15) !important;
            transition: all 0.2s ease !important;
            margin-top: 10px !important;
        }
        div.stDownloadButton > button:hover {
            background-color: #1f6b5c !important;
            color: #ffffff !important;
            box-shadow: 0 6px 10px rgba(0, 0, 0, 0.2) !important;
            transform: translateY(-1px);
        }
    </style>
""", unsafe_allow_html=True)

if 'pagina_atual' not in st.session_state:
    st.session_state.pagina_atual = 'home'

def navegar_para(pagina):
    st.session_state.pagina_atual = pagina


# ==========================================
# TELA INICIAL (MENU)
# ==========================================
if st.session_state.pagina_atual == 'home':
    st.markdown("""
        <style>
            div.stButton > button:first-child {
                font-size: 18px !important;
                font-weight: 600 !important;
                padding: 20px 25px !important;
                border-radius: 10px !important;
                border: none !important;
                box-shadow: 0 4px 6px rgba(0,0,0,0.05);
                transition: all 0.3s ease;
                width: 100%;
            }
            div.stColumn:nth-child(1) div.stButton > button {
                background-color: #d1ecf1 !important;
                color: #0c5460 !important;
            }
            div.stColumn:nth-child(1) div.stButton > button:hover {
                background-color: #bee5eb !important;
            }
            div.stColumn:nth-child(2) div.stButton > button {
                background-color: #d4edda !important;
                color: #155724 !important;
            }
            div.stColumn:nth-child(2) div.stButton > button:hover {
                background-color: #c3e6cb !important;
            }
        </style>
        <div class="header-box">
            <h1>REPORTA ACADÊMICO</h1>
            <p>Plataforma de consolidação de relatórios acadêmicos.</p>
            <h3>PPGOLD - UFPR</h3>
        </div>
    """, unsafe_allow_html=True)

    st.write("### Selecione o Módulo de Extração:")
    st.write("")
    st.write("")
    
    col1, col2 = st.columns(2)
    with col1:
        st.button("Extração dados para relatório CAPES", on_click=navegar_para, args=('capes',), use_container_width=True)
    with col2:
        st.button("Extração dados PTTs", on_click=navegar_para, args=('ptts',), use_container_width=True)


# ==========================================
# TELA 1: RELATÓRIO CAPES (LATTES)
# ==========================================
elif st.session_state.pagina_atual == 'capes':
    st.markdown("""
        <style>
            .stApp {
                background-color: #d1ecf1;
            }
            div.stColumn:nth-child(1) div.stButton > button {
                background-color: rgba(255, 255, 255, 0.5) !important;
                color: #2a8b79 !important;
                border: 1px solid rgba(42, 139, 121, 0.3) !important;
                padding: 2px 10px !important;
                font-size: 16px !important;
                font-weight: 600 !important;
                border-radius: 6px !important;
                width: auto !important;
                min-height: 32px !important;
                height: 32px !important;
                line-height: 1 !important;
                box-shadow: none !important;
            }
            div.stColumn:nth-child(1) div.stButton > button:hover {
                background-color: rgba(255, 255, 255, 0.9) !important;
                border-color: #2a8b79 !important;
            }
        </style>
    """, unsafe_allow_html=True)

    col_nav, col_titulo, col_vazia = st.columns([1, 10, 1])
    with col_nav:
        if st.button("←", key="voltar_capes"):
            navegar_para('home')
            st.rerun()
    with col_titulo:
        st.markdown('<div class="small-header">REPORTA ACADÊMICO</div>', unsafe_allow_html=True)

    st.divider()
    st.subheader("Extração para Relatório CAPES")
    
    with st.sidebar:
        st.write("Insira as planilhas oficiais.")
        base_abdc_raw = st.file_uploader("Upload ABDC", type=['xlsx', 'csv', 'xls'])
        base_abs_raw = st.file_uploader("Upload ABS / AJG", type=['xlsx', 'csv', 'xls'])
        base_jcr_raw = st.file_uploader("Upload JCR (Q1 a Q4)", type=['csv', 'xlsx', 'xls'], accept_multiple_files=True)
        base_sjr_raw = st.file_uploader("Upload SJR", type=['csv', 'xlsx', 'xls'])
        base_spell_raw = st.file_uploader("Upload SPELL", type=['xls', 'xlsx', 'csv'])
        base_scielo_raw = st.file_uploader("Upload SciELO", type=['csv', 'xlsx', 'xls'])
        
        base_abdc = validar_e_carregar_planilha(base_abdc_raw, "ABDC")
        base_abs = validar_e_carregar_planilha(base_abs_raw, "ABS")
        base_jcr = [validar_e_carregar_planilha(f, "JCR") for f in base_jcr_raw] if base_jcr_raw else []
        base_sjr = validar_e_carregar_planilha(base_sjr_raw, "SJR")
        base_spell = validar_e_carregar_planilha(base_spell_raw, "SPELL")
        base_scielo = validar_e_carregar_planilha(base_scielo_raw, "SciELO")
        
        bases_carregadas = []
        if base_abdc: bases_carregadas.append("ABDC")
        if base_abs: bases_carregadas.append("ABS")
        if base_jcr and any(base_jcr): bases_carregadas.append("JCR")
        if base_sjr: bases_carregadas.append("SJR")
        if base_spell: bases_carregadas.append("SPELL")
        if base_scielo: bases_carregadas.append("SciELO")

    ano_selecionado = st.number_input("Ano da Extração:", min_value=2000, max_value=2050, value=datetime.date.today().year, step=1)
    arquivos_xml = st.file_uploader("Insira os currículos Lattes", type=['xml'], accept_multiple_files=True)
    
    tem_arquivos = arquivos_xml is not None and len(arquivos_xml) > 0
    if tem_arquivos:
        st.markdown("""
            <style>
                div.stButton > button[kind="primary"] {
                    background-color: #dc3545 !important;
                    color: white !important;
                }
                div.stButton > button[kind="primary"]:hover {
                    background-color: #c82333 !important;
                }
            </style>
        """, unsafe_allow_html=True)
    
    if st.button("Processar Currículos", type="primary", disabled=not arquivos_xml):
        dict_bases = {'abdc': {}, 'abs': {}, 'jcr': {}, 'sjr': {}, 'spell': {}, 'scielo': set()}
        
        with st.spinner("Lendo Lattes e cruzando bases..."):
            try:
                if base_abdc:
                    df = pd.read_excel(base_abdc, skiprows=7)
                    for _, r in df.iterrows():
                        issn_print = str(r.get('ISSN')).replace('\t', '').strip()
                        issn_online = str(r.get('ISSNOnline')).replace('\t', '').strip()
                        rating = str(r.get('2025 rating')).replace('\t', '').strip()
                        if issn_print and issn_print != 'nan': dict_bases['abdc'][issn_print] = rating
                        if issn_online and issn_online != 'nan': dict_bases['abdc'][issn_online] = rating
                if base_abs:
                    df = pd.read_excel(base_abs) 
                    for _, r in df.iterrows():
                        if pd.notna(r.get('ISSN')): dict_bases['abs'][str(r['ISSN']).strip()] = str(r.get('Rating'))
                if base_jcr:
                    for f in base_jcr:
                        if f:
                            df = pd.read_csv(f, skiprows=2, on_bad_lines='skip')
                            for _, r in df.iterrows():
                                if pd.notna(r.get('ISSN')): dict_bases['jcr'][str(r['ISSN']).strip()] = r.get('JIF Quartile')
                                if pd.notna(r.get('eISSN')): dict_bases['jcr'][str(r['eISSN']).strip()] = r.get('JIF Quartile')
                if base_sjr:
                    df = pd.read_csv(base_sjr, sep=';', low_memory=False)
                    if len(df.columns) == 1: df = pd.read_csv(base_sjr, sep=',', low_memory=False)
                    for _, r in df.iterrows():
                        issns = str(r.get('Issn', '')).split(',')
                        for i in issns:
                            i = i.strip()
                            if len(i) == 8: i = f"{i[:4]}-{i[4:]}"
                            dict_bases['sjr'][i] = r.get('SJR Best Quartile')
                if base_scielo:
                    df = pd.read_csv(base_scielo)
                    for _, r in df.iterrows():
                        if pd.notna(r.get('Title')): dict_bases['scielo'].add(normalizar_nome(r['Title']))
                if base_spell:
                    try: df = pd.read_csv(base_spell, sep='\t', decimal=',')
                    except: df = pd.read_excel(base_spell)
                    col_impacto = [c for c in df.columns if 'Impacto 5 Anos' in c]
                    if col_impacto:
                        df['Impacto_Num'] = pd.to_numeric(df[col_impacto[0]].astype(str).str.replace(',', '.'), errors='coerce')
                        df = df.dropna(subset=['Impacto_Num']).sort_values('Impacto_Num', ascending=False)
                        quantis = [df['Impacto_Num'].quantile(q/10.0) for q in range(9, 0, -1)]
                        for _, r in df.iterrows():
                            nome = normalizar_nome(r.get('Periódico'))
                            imp_val = r['Impacto_Num']
                            rank_str = "Ausente"
                            if imp_val >= quantis[0]: rank_str = 'Top 10% (D1)'
                            elif imp_val >= quantis[1]: rank_str = 'D2 (11-20%)'
                            elif imp_val >= quantis[2]: rank_str = 'D3 (21-30%)'
                            elif imp_val >= quantis[3]: rank_str = 'D4 (31-40%)'
                            elif imp_val >= quantis[4]: rank_str = 'D5 (41-50%)'
                            elif imp_val >= quantis[5]: rank_str = 'D6 (51-60%)'
                            elif imp_val >= quantis[6]: rank_str = 'D7 (61-70%)'
                            elif imp_val >= quantis[7]: rank_str = 'D8 (71-80%)'
                            elif imp_val >= quantis[8]: rank_str = 'D9 (81-90%)'
                            else: rank_str = 'D10 (91-100%)'
                            dict_bases['spell'][nome] = {'rank': rank_str}
            except Exception as e:
                st.error(f"Erro ao ler uma das bases: {e}")

            todos_artigos = []
            
            for arquivo in arquivos_xml:
                arquivo.seek(0)
                todos_artigos.extend(extrair_dados_lattes(arquivo))
            
            tabela_artigos_final = pd.DataFrame()
            
            if todos_artigos:
                tb_artigos = pd.DataFrame(todos_artigos)
                tb_artigos['Ano_Num'] = pd.to_numeric(tb_artigos['Ano'], errors='coerce')
                tb_artigos = tb_artigos[tb_artigos['Ano_Num'] == ano_selecionado].copy()
                if not tb_artigos.empty:
                    tb_artigos['Titulo_Comparacao'] = tb_artigos['Titulo'].apply(normalizar_nome)
                    tabela_agrupada = tb_artigos.groupby('Titulo_Comparacao').agg({
                        'Professor': lambda x: ' | '.join(sorted(list(set(x)))),
                        'Titulo': 'first', 'Ano': 'first', 'Revista': 'first', 'Revista_Norm': 'first',
                        'ISSN': 'first', 'Abrangência': 'first'
                    }).reset_index(drop=True)
                    tabela_agrupada.rename(columns={'Professor': 'Autores (PPGOLD)'}, inplace=True)
                    tabela_agrupada['Qtd. Autores'] = tabela_agrupada['Autores (PPGOLD)'].apply(lambda x: len(x.split(' | ')))
                    tabela_agrupada[['Abrangência', 'ABDC', 'ABS', 'JCR', 'SJR', 'SciELO', 'SPELL (Classificação)', 'Conceito Área 27', 'Pontos']] = tabela_agrupada.apply(lambda row: calcular_nota_area27(row, dict_bases, bases_carregadas), axis=1)
                    tabela_artigos_final = tabela_agrupada[['Autores (PPGOLD)', 'Qtd. Autores', 'Titulo', 'Ano', 'Revista', 'ISSN', 'Abrangência', 'ABDC', 'ABS', 'JCR', 'SJR', 'SciELO', 'SPELL (Classificação)', 'Conceito Área 27', 'Pontos']]

            st.markdown("""
                <div style="background-color: #d4edda; border-left: 6px solid #28a745; padding: 14px 20px; border-radius: 4px; margin-bottom: 20px;">
                    <span style="color: #155724; font-size: 18px; font-weight: bold;">Processamento Lattes concluído com sucesso!</span>
                </div>
            """, unsafe_allow_html=True)
            
            st.subheader("Resultado da Extração")
            if not tabela_artigos_final.empty: 
                st.dataframe(tabela_artigos_final, use_container_width=True, hide_index=True)
                
                excel_bytes = gerar_excel_capes_estilizado(tabela_artigos_final, ano_selecionado)
                st.download_button(
                    "Baixar", 
                    data=excel_bytes, 
                    file_name=f"Relatorio_Lattes_CAPES_{ano_selecionado}.xlsx", 
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            else: 
                st.info("Nenhum artigo processado para o ano selecionado.")


# ==========================================
# TELA 2: EXTRAÇÃO DE PTTs (PDF / WORD)
# ==========================================
elif st.session_state.pagina_atual == 'ptts':
    st.markdown("""
        <style>
            .stApp {
                background-color: #d4edda;
            }
            div.stColumn:nth-child(1) div.stButton > button {
                background-color: rgba(255, 255, 255, 0.5) !important;
                color: #2a8b79 !important;
                border: 1px solid rgba(42, 139, 121, 0.3) !important;
                padding: 2px 10px !important;
                font-size: 16px !important;
                font-weight: 600 !important;
                border-radius: 6px !important;
                width: auto !important;
                min-height: 32px !important;
                height: 32px !important;
                line-height: 1 !important;
                box-shadow: none !important;
            }
            div.stColumn:nth-child(1) div.stButton > button:hover {
                background-color: rgba(255, 255, 255, 0.9) !important;
                border-color: #2a8b79 !important;
            }
            [data-testid="stFileUploaderDropzoneInstructions"] small {
                font-size: 0 !important;
            }
            [data-testid="stFileUploaderDropzoneInstructions"] small::after {
                content: "200MB per file • PDF, Word (DOCX, DOC)";
                font-size: 14px !important;
                color: rgba(49, 51, 63, 0.6);
            }
        </style>
    """, unsafe_allow_html=True)

    col_nav, col_titulo, col_vazia = st.columns([1, 10, 1])
    with col_nav:
        if st.button("←", key="voltar_ptts"):
            navegar_para('home')
            st.rerun()
    with col_titulo:
        st.markdown('<div class="small-header">REPORTA ACADÊMICO</div>', unsafe_allow_html=True)

    st.divider()
    st.subheader("Extração de Fichas de Avaliação PTT")
    st.write("Adicione as Fichas de Avaliação de Prática Profissional (Anexo 2) em formato PDF ou Word, o sistema identificará automaticamente todos os campos relevantes: identificação, resumo, tipo de trabalho e as respostas assinaladas nos critérios.")
    
    arquivos_ptt = st.file_uploader("Arraste as Fichas PTT.", type=['pdf', 'docx', 'doc'], accept_multiple_files=True)
    
    tem_arquivos_ptt = arquivos_ptt is not None and len(arquivos_ptt) > 0
    if tem_arquivos_ptt:
        st.markdown("""
            <style>
                div.stButton > button[kind="primary"] {
                    background-color: #dc3545 !important;
                    color: white !important;
                }
                div.stButton > button[kind="primary"]:hover {
                    background-color: #c82333 !important;
                }
            </style>
        """, unsafe_allow_html=True)
    
    if st.button("Extrair Dados das Fichas", type="primary", disabled=not arquivos_ptt):
        with st.spinner("Extraindo informações dos arquivos..."):
            todos_ptts_pdf = []
            for arq in arquivos_ptt:
                todos_ptts_pdf.append(extrair_dados_arquivo_ptt(arq))
                
            tabela_ptts_pdf_final = pd.DataFrame(todos_ptts_pdf)
            
            st.markdown(f"""
                <div style="background-color: #c3e6cb; border-left: 6px solid #28a745; padding: 14px 20px; border-radius: 4px; margin-bottom: 20px;">
                    <span style="color: #155724; font-size: 18px; font-weight: bold;">Extração concluída! {len(tabela_ptts_pdf_final)} arquivo(s) processado(s).</span>
                </div>
            """, unsafe_allow_html=True)

            st.subheader("Resultado da Extração")
            st.dataframe(tabela_ptts_pdf_final, use_container_width=True, hide_index=True)
            
            excel_ptt_bytes = gerar_excel_ptts_estilizado(tabela_ptts_pdf_final)
                
            st.download_button(
                "Baixar", 
                data=excel_ptt_bytes, 
                file_name="Relatorio_Geral_Fichas_PTT.xlsx", 
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )